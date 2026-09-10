# -*- coding: utf-8 -*-
r"""El Excel de partidas que la web exporta y el mapa local importa.

POR QUE EXISTE
--------------
La web sondea AA2000 cada 5 minutos, 24/7, y tiene el acumulado bueno. El mapa que corre
en la terminal no puede hacer lo mismo: su regla 2 es arrancar sin internet, y en esa PC
el feed puede estar bloqueado por la red (medido: 403 URLBlocked).

Entonces el dato viaja en un archivo. Cada tanto se baja este Excel de la web y se pega en
la biblioteca de SharePoint; la terminal lo lee de la carpeta **sincronizada en disco** y
actualiza sus datos. Tiempo real en la web, por lotes en la terminal.

**Eso NO rompe la regla 2**, y el motivo importa: una biblioteca sincronizada es una
carpeta local. El mapa abre un archivo del disco; la sincronizacion la hace el cliente de
OneDrive, que es otro programa. Es el mismo caso que `datos/semilla.db`.

UN SOLO MODULO PARA ESCRIBIR Y PARA LEER
----------------------------------------
Este archivo esta en los DOS repos y define el formato una sola vez. Con un escritor y un
lector separados, un dia uno agrega una columna, el otro no la espera, y el archivo pasa a
significar cosas distintas segun quien lo abrio. Ya paso en este proyecto con la planilla
de proveedores, y por eso `cargar_excel` es compartido.

QUE LLEVA: LAS PARTIDAS CRUDAS, NO LOS M3
-----------------------------------------
Es la decision de fondo. El Excel podria traer el consumo ya calculado -- seria mas chico
y mas facil de leer a ojo -- y esta mal por dos razones:

  1. **Habria dos calculos de combustible.** Todo el repo tiene uno: `avion_model`. Con el
     m3 escrito en el archivo, el numero de la terminal saldria de la web y el del mapa de
     su propio modelo, y un dia difieren sin que nadie sepa cual creer.
  2. **El historico no se podria recalcular.** Si manana se calibra la flota o se corrige
     una curva, con las partidas crudas todo el pasado se recalcula solo; con los m3
     escritos quedan congelados los de la version vieja.

El costo es un archivo mas grande, y no es un costo: ~1.000 filas por dia.

FORMATO
-------
Dos hojas, y las dos hacen falta:

  `partidas`  una fila por partida, con las columnas de `vuelo_oficial` que el mapa
              consume. La primera fila son los encabezados.
  `meta`      clave/valor: cuando se genero, hasta que momento tiene datos, cuantas filas
              y de que version del formato es.

**La hoja `meta` es la que evita el error mas grave de todo esto**: que la terminal muestre
datos de hace una semana como si fueran de hoy. Un dato que se presenta como actual y esta
congelado es peor que uno viejo con fecha, y ya lo resolvimos dos veces en este proyecto
(el JSON de TimesFM y `proveedores.json`). Sin `meta` no hay forma de saberlo desde el
archivo.
"""
import io
import os

# Version del formato. Si algun dia cambian las columnas, esto sube y el lector puede
# decir "este archivo es de un formato mas nuevo, actualiza el mapa" en vez de leer mal
# en silencio.
FORMATO = 1

HOJA_PARTIDAS = 'partidas'
HOJA_META = 'meta'

# UN NOMBRE FIJO QUE SE PISA, desde el 2026-09-10. Actualizar los datos de la terminal
# es reemplazar este archivo y nada mas: no hay que borrar el anterior ni acordarse de
# ningun formato de fecha. La fecha de los datos NO se pierde por eso -- viaja en la
# hoja `meta` (`hasta`), que es de donde sale el cartel del filtro, en ambar cuando el
# archivo quedo viejo. Ponerla ademas en el nombre daba dos fuentes para el mismo dato.
#
# Va en MAYUSCULAS y con espacio porque asi se llama en la biblioteca compartida, y el
# nombre que ve la gente y el que busca el codigo tienen que ser el mismo. La busqueda
# igual es case-insensitive y acepta el nombre viejo (ver PREFIJOS en importar_ms.py).
NOMBRE_ARCHIVO = 'BAJADA MAPA.xlsx'

# Las columnas que viajan, en orden. Son las que `msrtic` consume mas las que
# `aa2000.guardar()` no degrada: no va todo `vuelo_oficial` porque la mitad son campos de
# la pantalla de AA2000 (cinta, puerta, sector) que el mapa no usa para nada.
COLUMNAS = ('id', 'aeropuerto', 'movimiento', 'numero', 'aerolinea_id', 'aerolinea',
            'otro_aeropuerto', 'destino_nombre', 'programada', 'programada_epoch',
            'real', 'real_epoch', 'estado', 'matricula', 'pasajeros', 'cuerpo',
            'tipo_vuelo', 'primera_vez', 'ultima_vez', 'veces_visto')

# Sin estas cuatro no se puede armar una ruta ni deduplicar, asi que un archivo que no las
# traiga se rechaza entero en vez de importar filas a medias.
IMPRESCINDIBLES = ('id', 'aeropuerto', 'movimiento', 'otro_aeropuerto')


def _valor_excel(v):
    """Lo que openpyxl puede escribir sin quejarse."""
    if v is None or isinstance(v, (int, float, str, bool)):
        return v
    return str(v)


def escribir(filas, meta, destino=None):
    """Arma el Excel. `destino` es una ruta, o None para devolver los bytes.

    `filas` son dicts (una partida cada uno) y `meta` un dict clave/valor.
    """
    from openpyxl import Workbook

    wb = Workbook()
    hp = wb.active
    hp.title = HOJA_PARTIDAS
    hp.append(list(COLUMNAS))
    for f in filas:
        hp.append([_valor_excel(f.get(c)) for c in COLUMNAS])

    hm = wb.create_sheet(HOJA_META)
    hm.append(['clave', 'valor'])
    for k, v in sorted(meta.items()):
        hm.append([k, _valor_excel(v)])

    if destino:
        wb.save(destino)
        return destino
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def leer(origen):
    """(filas, meta) del Excel. Lanza ValueError con un motivo legible si no sirve.

    `origen` es una ruta o un objeto de archivo. Los errores se explican en castellano y
    en una linea porque el destinatario es alguien mirando una terminal, no un log.
    """
    from openpyxl import load_workbook

    try:
        # read_only para no cargar el libro entero en memoria, y data_only para que una
        # celda con formula traiga el VALOR: si el libro tiene formulas sin recalcular el
        # valor queda None, y eso se cuenta como fila incompleta en vez de entrar como
        # basura.
        wb = load_workbook(origen, read_only=True, data_only=True)
    except Exception as e:
        raise ValueError('no se pudo abrir el Excel (%s)' % type(e).__name__)

    try:
        if HOJA_PARTIDAS not in wb.sheetnames:
            raise ValueError("el Excel no tiene la hoja '%s'. Hojas: %s"
                             % (HOJA_PARTIDAS, ', '.join(wb.sheetnames) or 'ninguna'))

        meta = {}
        if HOJA_META in wb.sheetnames:
            for fila in wb[HOJA_META].iter_rows(min_row=2, values_only=True):
                if fila and fila[0]:
                    meta[str(fila[0])] = fila[1] if len(fila) > 1 else None

        it = wb[HOJA_PARTIDAS].iter_rows(values_only=True)
        try:
            encabezados = [str(c).strip() if c is not None else '' for c in next(it)]
        except StopIteration:
            raise ValueError('la hoja de partidas esta vacia')

        # Se mapea POR NOMBRE y no por posicion: si alguien reordena las columnas en el
        # Excel, leer por posicion mete la matricula en el campo de la aerolinea sin que
        # nada falle. Las columnas que falten quedan en None, que es lo que el resto del
        # codigo ya sabe manejar.
        faltan = [c for c in IMPRESCINDIBLES if c not in encabezados]
        if faltan:
            raise ValueError('al Excel le faltan columnas imprescindibles: %s'
                             % ', '.join(faltan))
        idx = {c: encabezados.index(c) for c in COLUMNAS if c in encabezados}

        filas = []
        for fila in it:
            if not fila:
                continue
            d = {c: fila[i] if i < len(fila) else None for c, i in idx.items()}
            if not d.get('id'):
                continue          # sin id no se puede deduplicar: no entra
            filas.append(d)
        if not filas:
            raise ValueError('el Excel no tiene ninguna partida con id')
        return filas, meta
    finally:
        wb.close()


def importar(filas, conn):
    """Mete las filas en `vuelo_oficial` sin degradar lo que ya esta.

    Devuelve {'nuevas', 'completadas', 'sin_cambio'}.

    Por que no un simple INSERT OR REPLACE: la base local puede tener la hora real de un
    despegue que el archivo no trajo (o al reves), y reemplazar la fila entera la
    perderia. Esa hora vive unas horas en el feed de AA2000 y despues no vuelve nunca, asi
    que perderla es definitivo.
    """
    cols_base = {r[1] for r in conn.execute('pragma table_info(vuelo_oficial)')}
    r = {'nuevas': 0, 'completadas': 0, 'sin_cambio': 0}

    for f in filas:
        ident = f.get('id')
        cur = conn.execute('select * from vuelo_oficial where id=?', (ident,))
        previa = cur.fetchone()
        usables = {k: v for k, v in f.items() if k in cols_base and v not in (None, '')}
        if not usables:
            continue

        if previa is None:
            cols = list(usables)
            conn.execute('insert into vuelo_oficial (%s) values (%s)'
                         % (','.join('"%s"' % c for c in cols), ','.join('?' * len(cols))),
                         [usables[c] for c in cols])
            r['nuevas'] += 1
            continue

        actual = dict(zip([d[0] for d in cur.description], previa))
        # Solo se completa lo que falta. Un valor presente no se toca ni siquiera si el
        # archivo trae otro: el archivo puede ser mas viejo que la base.
        cambios = {k: v for k, v in usables.items() if actual.get(k) in (None, '')}
        if not cambios:
            r['sin_cambio'] += 1
            continue
        conn.execute('update vuelo_oficial set %s where id=?'
                     % ','.join('"%s"=?' % c for c in cambios),
                     list(cambios.values()) + [ident])
        r['completadas'] += 1

    conn.commit()
    return r


def ruta_sharepoint_por_defecto():
    r"""Donde se busca el Excel en la PC de la terminal, si no se configura otra cosa.

    La biblioteca sincronizada, informada el 2026-09-08:

        %USERPROFILE%\OneDrive - YPF\GR - Aviación - Comercial - SEGUIMIENTO PLAN DIARIO

    **El usuario NO se hardcodea** (en esa maquina es YBL2961 y aca es otro), asi que sale
    de %USERPROFILE% y el resto va relativo. El nombre lleva espacios, guiones y un
    acento, y el de la carpeta de OneDrive incluye el tenant ("OneDrive - YPF"), que
    cambia si cambia la cuenta -- de ahi que todo se arme con os.path.join y que
    MS_RTIC_EXCEL pueda pisarlo entero.
    """
    return os.path.join(
        os.path.expanduser('~'),
        'OneDrive - YPF',
        'GR - Aviación - Comercial - SEGUIMIENTO PLAN DIARIO')
