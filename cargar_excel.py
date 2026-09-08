"""Lee la planilla de proveedores de YPF y la convierte a proveedores.json.

POR DEFECTO NO ESCRIBE NADA. Muestra que hojas y columnas encontro, como
interpreto cada una y todo lo que le resulto raro; recien con --escribir genera
el JSON. No es cautela decorativa: el 2026-09-07 este proyecto tuvo dos bugs de
lectura de campos en un dia -- `id_arpt` que no filtraba y `arpt` que era el
origen y no el destino -- y los dos eran SILENCIOSOS. Una columna mal
interpretada aca hace exactamente lo mismo con el proveedor, y el resultado se
ve perfectamente normal en pantalla.

COMO SE INTERPRETA LA PLANILLA:

    PARTIDA    origen IATA de 3 letras: DONDE CARGA el avion
    ARRIBO     destino IATA de 3 letras
    PROVEEDOR  quien abastece esa partida (YPF, Axion, Raizen, ...)

Una fila es una ruta DIRIGIDA. La llave es (PARTIDA, ARRIBO) y no el par de
ciudades: AEP->BRC y BRC->AEP son dos rutas distintas con proveedores distintos.
Ver proveedores.py.

Los encabezados se buscan POR NOMBRE y no por posicion. Si la planilla cambia el
orden de las columnas, un lector por posicion sigue funcionando y devuelve el
campo equivocado sin avisar: es el bug de `arpt` otra vez.

Uso:
  python cargar_excel.py planilla.xlsx              revisar, sin escribir
  python cargar_excel.py planilla.xlsx --escribir   convertir a proveedores.json
"""
from __future__ import annotations

import argparse
import json
import sys
from datetime import date
from pathlib import Path

from proveedores import COMPETITIVOS

# Nombres aceptados para cada columna, en minusculas y sin espacios. Se aceptan
# variantes porque la planilla la escribe una persona, y "ORIGEN"/"DESTINO" es
# tan probable como "PARTIDA"/"ARRIBO".
ALIAS = {
    "partida": {"partida", "origen", "desde", "aeropuertopartida", "salida"},
    "arribo": {"arribo", "destino", "hasta", "aeropuertoarribo", "llegada"},
    "proveedor": {"proveedor", "proveedorcombustible", "abastecedor",
                  "compania", "empresa"},
}


def _norm(v) -> str:
    return "".join(str(v or "").split()).lower()


def _iata(v) -> str:
    return str(v or "").strip().upper()


def _encabezado(filas):
    """Busca la fila de encabezados: (indice, {campo: columna}) o None.

    Se recorren las primeras 20 filas porque una planilla real suele tener un
    titulo o una fila en blanco arriba. Hacen falta las TRES columnas: con dos
    no se puede afirmar que esa fila sea el encabezado.
    """
    for i, fila in enumerate(filas[:20]):
        pos = {}
        for j, celda in enumerate(fila):
            n = _norm(celda)
            for campo, nombres in ALIAS.items():
                if n in nombres and campo not in pos:
                    pos[campo] = j
        if len(pos) == 3:
            return i, pos
    return None


def leer(xlsx: Path) -> dict:
    """Todo lo que hay en el libro, con lo raro contado y no descartado."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise SystemExit("falta openpyxl:  pip install openpyxl")

    # data_only=True para que una celda con formula traiga el VALOR y no el
    # texto de la formula. Si el libro tiene formulas sin recalcular el valor
    # queda None, y eso se cuenta como fila incompleta en vez de entrar como
    # ruta vacia.
    wb = load_workbook(xlsx, data_only=True, read_only=True)
    r = {"archivo": str(xlsx), "hojas": [], "rutas": {},
         "conflictos": [], "repetidas": [], "auto": [],
         "origen_no_competitivo": [], "incompletas": 0, "filas_leidas": 0}
    for hoja in wb.worksheets:
        filas = [tuple(f) for f in hoja.iter_rows(values_only=True)]
        enc = _encabezado(filas)
        info = {"nombre": hoja.title, "filas_totales": len(filas),
                "encabezado": None, "columnas": None, "filas_datos": 0,
                "proveedores": {}, "motivo": None}
        if enc is None:
            # Una hoja sin las tres columnas no se lee, y SE DICE. Que el libro
            # traiga una hoja de notas es normal; que traiga las rutas de EZE en
            # una hoja que no supimos leer no lo es, y la unica forma de
            # distinguir los dos casos es nombrarla.
            info["motivo"] = "no se encontraron las tres columnas"
            r["hojas"].append(info)
            continue
        i0, pos = enc
        info["encabezado"] = i0 + 1              # 1-based, como lo ve Excel
        info["columnas"] = {k: v + 1 for k, v in pos.items()}
        for n, fila in enumerate(filas[i0 + 1:], start=i0 + 2):
            def celda(campo, _fila=fila):
                j = pos[campo]
                return _fila[j] if j < len(_fila) else None
            o, dst = _iata(celda("partida")), _iata(celda("arribo"))
            prov = str(celda("proveedor") or "").strip()
            r["filas_leidas"] += 1
            if not (o and dst and prov):
                r["incompletas"] += 1
                continue
            info["filas_datos"] += 1
            info["proveedores"][prov] = info["proveedores"].get(prov, 0) + 1
            ref = f"{hoja.title}!{n}"
            if o == dst:
                # UNA RUTA CONTRA SI MISMA ES UN ERROR DE LA PLANILLA, no un
                # caso raro a interpretar. Confirmado por Fran el 2026-09-07
                # sobre la fila AEP-AEP de su Excel: "AEP AEP NO EXISTE".
                #
                # Se anota y NO se carga. Y se reporta fuerte, porque es la
                # misma firma que tuvo el bug de `arpt` -- cuando ese campo se
                # leia como destino, TODAS las rutas salian AEP->AEP-- asi que
                # ver esto de golpe en muchas filas no dice "la planilla tiene
                # una fila de mas" sino "el lector esta tomando la columna
                # equivocada".
                r["auto"].append({"ruta": f"{o}-{dst}", "proveedor": prov,
                                  "fila": ref})
                continue
            if o not in COMPETITIVOS:
                # El interior ya es YPF por la regla de negocio, sin planilla.
                # Una fila asi es REDUNDANTE si dice YPF y es una CONTRADICCION
                # si dice otra cosa. Se carga igual -- la planilla es la fuente
                # y quien manda es YPF -- pero queda dicho, porque si no nadie
                # se entera de que la regla y la planilla no coinciden.
                r["origen_no_competitivo"].append(
                    {"ruta": f"{o}-{dst}", "proveedor": prov, "fila": ref,
                     "contradice": prov.upper() != "YPF"})
            k = (o, dst)
            if k in r["rutas"]:
                if r["rutas"][k]["proveedor"].upper() != prov.upper():
                    # DOS PROVEEDORES PARA LA MISMA RUTA. No se elige uno: se
                    # reporta y la ruta queda sin declarar. Quedarse con el
                    # ultimo haria que el resultado dependa del orden de las
                    # filas de la planilla.
                    r["conflictos"].append(
                        {"ruta": f"{o}-{dst}",
                         "proveedores": [r["rutas"][k]["proveedor"], prov],
                         "filas": [r["rutas"][k]["fila"], ref]})
                else:
                    r["repetidas"].append({"ruta": f"{o}-{dst}", "fila": ref})
                continue
            r["rutas"][k] = {"proveedor": prov, "fila": ref}
        r["hojas"].append(info)
    wb.close()
    # Las rutas en conflicto se sacan DESPUES de recorrer todo el libro: si se
    # sacaran al detectarlas, una tercera fila con el mismo par la volveria a
    # cargar y el conflicto quedaria reportado pero igual aplicado.
    for c in r["conflictos"]:
        o, _, dst = c["ruta"].partition("-")
        r["rutas"].pop((o, dst), None)
    return r


def resumen_json(r: dict) -> dict:
    """Lo mismo que imprime informe(), pero serializable.

    Existe para que la pantalla de carga muestre EXACTAMENTE el cuadro que
    muestra la consola. Si la pagina armara su propio resumen, un dia diria algo
    distinto de lo que dice el .bat sobre la misma planilla, y no habria forma
    de saber cual mirar.

    Las llaves de `rutas` son tuplas y JSON no las tiene, asi que se aplanan a
    "ORIGEN-DESTINO", igual que en proveedores.json.
    """
    origenes: dict[str, int] = {}
    provs: dict[str, int] = {}
    for (o, _d), v in r["rutas"].items():
        origenes[o] = origenes.get(o, 0) + 1
        provs[v["proveedor"]] = provs.get(v["proveedor"], 0) + 1
    return {
        "archivo": Path(r["archivo"]).name,
        "hojas": r["hojas"],
        "filas_leidas": r["filas_leidas"],
        "n_rutas": len(r["rutas"]),
        "incompletas": r["incompletas"],
        "por_origen": dict(sorted(origenes.items())),
        "por_proveedor": dict(sorted(provs.items())),
        # Los aeropuertos competitivos SIN ninguna ruta: es la advertencia mas
        # importante del cuadro. Si falta EZE, toda partida de ahi queda sin
        # proveedor y el hueco parece del sistema cuando es de la planilla.
        "sin_ninguna_ruta": sorted(COMPETITIVOS - set(origenes)),
        "conflictos": r["conflictos"],
        "auto": r["auto"],
        "origen_no_competitivo": r["origen_no_competitivo"],
        "n_repetidas": len(r["repetidas"]),
    }


def escribir_json(r: dict, salida: Path, fuente: str | None = None) -> int:
    """Vuelca las rutas leidas a proveedores.json. Devuelve cuantas escribio.

    Un solo escritor para la consola y para la pantalla de carga: con dos, uno
    se queda sin un cambio -- el formato de la fecha, el nombre de la fuente --
    y el archivo pasa a significar cosas distintas segun quien lo escribio.
    """
    salida.write_text(json.dumps({
        "actualizado": date.today().isoformat(),
        "fuente": fuente or f"planilla de YPF: {Path(r['archivo']).name}",
        "rutas": {f"{o}-{d}": v["proveedor"]
                  for (o, d), v in sorted(r["rutas"].items())},
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    return len(r["rutas"])


def informe(r: dict) -> None:
    print(f"\nPlanilla: {r['archivo']}")
    print(f"\n  {len(r['hojas'])} hoja(s) en el libro:")
    for h in r["hojas"]:
        if h["encabezado"] is None:
            print(f"    [{h['nombre']}]  {h['filas_totales']} filas"
                  f"  -- NO SE LEYO: {h['motivo']}")
            continue
        cols = "  ".join(f"{k}=col {v}" for k, v in h["columnas"].items())
        print(f"    [{h['nombre']}]  encabezado en fila {h['encabezado']}"
              f"  ({cols})")
        prov = "  ".join(f"{k}:{v}" for k, v in sorted(h["proveedores"].items()))
        print(f"        {h['filas_datos']} rutas leidas   {prov}")

    print("\n  ASI SE INTERPRETO:")
    print("    PARTIDA   = origen, DONDE CARGA el avion")
    print("    ARRIBO    = destino")
    print("    PROVEEDOR = quien abastecio esa partida")
    print("    la llave es la ruta DIRIGIDA (PARTIDA, ARRIBO):"
          " AEP-BRC y BRC-AEP son distintas")

    print(f"\n  {len(r['rutas'])} rutas cargadas de {r['filas_leidas']}"
          f" filas leidas")
    origenes = {}
    provs = {}
    for (o, _dst), v in r["rutas"].items():
        origenes[o] = origenes.get(o, 0) + 1
        provs[v["proveedor"]] = provs.get(v["proveedor"], 0) + 1
    print("    por origen:     "
          + "  ".join(f"{k}:{v}" for k, v in sorted(origenes.items())))
    print("    por proveedor:  "
          + "  ".join(f"{k}:{v}" for k, v in sorted(provs.items())))
    falta = sorted(COMPETITIVOS - set(origenes))
    if falta:
        # Es la advertencia mas importante del informe: si la planilla no trae
        # EZE, toda partida de Ezeiza va a salir "sin proveedor" y el hueco
        # parece un problema del sistema cuando es de la planilla.
        print(f"    SIN NINGUNA RUTA: {', '.join(falta)}"
              f"  -- toda partida de ahi va a quedar SIN PROVEEDOR")

    def bloque(titulo, xs, linea):
        if not xs:
            return
        print(f"\n  {titulo} ({len(xs)}):")
        for x in xs[:15]:
            print(f"    {linea(x)}")
        if len(xs) > 15:
            print(f"    ... y {len(xs) - 15} mas")

    bloque("CONFLICTO: misma ruta con dos proveedores -- NO se cargaron",
           r["conflictos"],
           lambda c: f"{c['ruta']}: {' vs '.join(c['proveedores'])}"
                     f"   filas {', '.join(c['filas'])}")
    bloque("ERROR EN LA PLANILLA: una ruta contra si misma no existe"
           " -- NO se cargaron", r["auto"],
           lambda x: f"{x['ruta']} = {x['proveedor']}   fila {x['fila']}"
                     f"   -- borrar esa fila")
    if len(r["auto"]) > 3:
        # UNA fila asi es un error de tipeo; MUCHAS son otra cosa. Cuando el bug
        # de `arpt` leia el origen como destino, TODAS las rutas salian contra si
        # mismas: si esto aparece en cantidad, el sospechoso son las columnas y
        # no la planilla.
        print("    OJO: son muchas para ser un error de tipeo. Cuando este")
        print("    proyecto leyo mal el campo del destino, TODAS las rutas")
        print("    salieron asi. Revisa que PARTIDA y ARRIBO no esten cruzadas.")

    bloque("ORIGEN QUE NO ES AEP/EZE/COR: el interior ya es YPF por regla",
           r["origen_no_competitivo"],
           lambda x: f"{x['ruta']} = {x['proveedor']}   fila {x['fila']}"
                     + ("   CONTRADICE LA REGLA" if x["contradice"]
                        else "   (redundante)"))
    bloque("repetidas con el mismo proveedor (sin efecto)", r["repetidas"],
           lambda x: f"{x['ruta']}   fila {x['fila']}")
    if r["incompletas"]:
        print(f"\n  {r['incompletas']} filas con alguna de las tres celdas"
              f" vacia, no cargadas")


def elegir_planilla(carpeta: Path) -> Path | None:
    """El .xlsx MAS RECIENTE de la carpeta, o None si no hay ninguno.

    El mas reciente y no el primero por nombre: la planilla se recarga cada
    tanto, asi que van a convivir varias versiones en la carpeta. Tomar la
    primera alfabeticamente cargaria la vieja -- y sin decirlo, porque una
    planilla vieja se lee igual de bien que una nueva.

    Se ignoran los temporales de Excel (`~$...`), que existen mientras el
    archivo esta abierto y openpyxl no puede leer.
    """
    xs = [x for x in carpeta.glob("*.xlsx") if not x.name.startswith("~$")]
    if not xs:
        return None
    return max(xs, key=lambda x: x.stat().st_mtime)


def hay_que_recargar(xlsx: Path, salida: Path) -> bool:
    """Si la planilla es mas nueva que lo ya cargado.

    Por fecha de modificacion y no por contenido: alcanza para el caso real
    -- YPF manda una version nueva y hay que pisar-- y no obliga a leer el
    libro entero para decidir si conviene leerlo.
    """
    if not salida.exists():
        return True
    return xlsx.stat().st_mtime > salida.stat().st_mtime


def main() -> int:
    p = argparse.ArgumentParser(
        description=__doc__,
        formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("xlsx", nargs="?", help="la planilla de YPF")
    p.add_argument("--auto", action="store_true",
                   help="buscar la planilla mas reciente de esta carpeta y"
                        " revisarla solo si cambio. Sale con 10 si hay algo"
                        " que confirmar (lo usa el .bat)")
    p.add_argument("--escribir", action="store_true",
                   help="generar proveedores.json (sin esto solo se revisa)")
    p.add_argument("--salida", default="proveedores.json")
    a = p.parse_args()

    salida = Path(a.salida)
    if a.auto:
        # El .bat corre esto en cada arranque. Tiene que ser silencioso y salir
        # con 0 cuando no hay nada que hacer: preguntar por una planilla que no
        # cambio, en todas las corridas, termina en que nadie lee la pregunta.
        ruta = elegir_planilla(Path(__file__).parent)
        if ruta is None:
            print("  No hay ninguna planilla .xlsx en esta carpeta.")
            if not salida.exists():
                print("  Copia el Excel de YPF aca, o arrastralo sobre el .bat.")
            return 0
        if not hay_que_recargar(ruta, salida):
            print(f"  Planilla ya cargada y sin cambios: {ruta.name}")
            print(f"  Para volver a cargarla, guardala de nuevo o borra"
                  f" {salida.name}.")
            return 0
        print(f"  Planilla {'nueva' if salida.exists() else 'encontrada'}:"
              f" {ruta.name}")
        if salida.exists():
            print(f"  Es mas nueva que {salida.name}: si confirmas, LO PISA.")
    else:
        if not a.xlsx:
            p.error("falta la planilla (o usa --auto)")
        ruta = Path(a.xlsx)
    if not ruta.exists():
        print(f"no existe: {ruta}")
        return 1
    r = leer(ruta)
    informe(r)

    if not r["rutas"]:
        print("\nNO SE CARGO NINGUNA RUTA. No se escribe nada.")
        return 1
    if not a.escribir:
        print("\n  --- NO SE ESCRIBIO NADA ---")
        if a.auto:
            # 10 y no 0: el .bat lo lee para saber si tiene que preguntar. Un 0
            # aca lo haria seguir de largo y la planilla no se cargaria nunca.
            return 10
        print("  Revisa el cuadro de arriba. Si la interpretacion es correcta:")
        print(f'      python cargar_excel.py "{ruta}" --escribir')
        return 0

    n = escribir_json(r, salida)
    print(f"\n  escrito {salida}  ({n} rutas)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
